"""
app.py - ingest + report service.

Ingest paths (all end up in raw_messages, de-duplicated by Telegram message id):
  POST /ingest            <- your Apps Script (recommended; header X-Token)
  POST /webhook/telegram  <- a Telegram bot webhook (only for a NEW bot, see README)
  import_telegram_export.py  <- backfill from a Telegram Desktop JSON export
Views:
  GET /?period=2026-09 | 2026-W40        HTML report
  GET /report.csv?period=...  /report.xlsx?period=...  /api/report?period=...
  GET /review                             messages the parser could not understand
  POST /import-export                     backfill from a Telegram Desktop JSON export
  POST /rebuild                           re-parse everything after fixing a pattern
  POST /api/manual-event                  manual override (symbol, date, kind, pct)
"""
import csv, html, io, os
from datetime import datetime, timedelta, timezone
import yaml
from fastapi import FastAPI, Header, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response

import db
from engine import build_report, label_for

CFG = yaml.safe_load(open(os.path.join(os.path.dirname(__file__), "config.yaml")))
OFFSET = timedelta(minutes=CFG.get("timezone_offset_minutes", 330))
INGEST_TOKEN = os.environ.get("INGEST_TOKEN", "")
VIEW_KEY = os.environ.get("VIEW_KEY", "")
TG_SECRET = os.environ.get("TG_WEBHOOK_SECRET", "")

app = FastAPI(title="SuperSignal performance tracker")
db.init()


def ist_iso(unix: int) -> str:
    return (datetime.fromtimestamp(unix, tz=timezone.utc) + OFFSET).strftime("%Y-%m-%dT%H:%M:%S")


def need_ingest(tok):
    if not INGEST_TOKEN or tok != INGEST_TOKEN:
        raise HTTPException(401, "bad token")


def need_view(key):
    if VIEW_KEY and key != VIEW_KEY:
        raise HTTPException(401, "add ?key=...")


# ---------------- ingest ----------------
@app.post("/ingest")
async def ingest(req: Request, x_token: str = Header(default="")):
    need_ingest(x_token)
    b = await req.json()
    db.add_raw("telegram", int(b["message_id"]), ist_iso(int(b["date"])), b["text"])
    db.rebuild()
    return {"ok": True}


@app.post("/webhook/telegram")
async def webhook(req: Request, x_telegram_bot_api_secret_token: str = Header(default="")):
    if not TG_SECRET or x_telegram_bot_api_secret_token != TG_SECRET:
        raise HTTPException(401, "bad secret")
    u = await req.json()
    m = u.get("channel_post") or u.get("edited_channel_post")
    if m and m["chat"]["id"] == CFG["telegram"]["test_channel_id"]:
        text = m.get("text") or m.get("caption") or ""
        if text:
            db.add_raw("telegram", m["message_id"], ist_iso(m["date"]), text)
            db.rebuild()
    return {"ok": True}


@app.post("/import-export")
async def import_export(req: Request, x_token: str = Header(default="")):
    """Upload a Telegram Desktop result.json (test channel only) to backfill history."""
    need_ingest(x_token)
    from import_telegram_export import import_data
    return {"ok": True, "imported": import_data(await req.json())}


@app.post("/rebuild")
def rebuild(x_token: str = Header(default="")):
    need_ingest(x_token)
    db.rebuild()
    return {"ok": True}


@app.post("/api/manual-event")
async def manual_event(req: Request, x_token: str = Header(default="")):
    need_ingest(x_token)
    b = await req.json()
    assert b["kind"] in {"TARGET", "SL_HIT", "TSL_HIT", "INVALID", "MANUAL"}
    db.add_manual_event(b["symbol"], b["date"], b["kind"], b.get("pct"))
    return {"ok": True}


# ---------------- reports ----------------
def current_period(kind="month"):
    return label_for(kind, (datetime.now(timezone.utc) + OFFSET).date())


def report_for(period):
    sigs, evs = db.load()
    return build_report(sigs, evs, period)


COLS = ["signal_date", "symbol", "tag", "entry", "sl", "target", "status", "l1", "l2", "kind", "origin", "note"]
HEAD = ["Date", "Stock", "Tag", "Entry", "SL", "Target", "Status", "L1 %", "L2 %", "Type", "Signalled in", "Note"]


def table(rows, cols=COLS, head=HEAD):
    if not rows:
        return "<p class=muted>None.</p>"
    th = "".join(f"<th>{h}</th>" for h in head)
    body = ""
    for r in rows:
        tds = ""
        for c in cols:
            v = r.get(c, "")
            v = "" if v is None else v
            cls = ""
            if c in ("l1", "l2") and isinstance(v, (int, float)):
                cls = "pos" if v > 0 else "neg" if v < 0 else ""
            tds += f"<td class='{cls}'>{html.escape(str(v))}</td>"
        body += f"<tr>{tds}</tr>"
    return f"<div class=scroll><table><tr>{th}</tr>{body}</table></div>"


def stat_block(title, s):
    return (f"<div class=stat><h3>{title}</h3><div class=big>{s['win_rate']}%</div>"
            f"<div>{s['winners']} winners / {s['losers']} losers of {s['counted']}</div>"
            f"<div>Total {s['total']:+}% &middot; Avg {s['avg']:+}%</div></div>")


PAGE = """<!doctype html><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1">
<title>SuperSignal performance {period}</title><style>
:root{{--bg:#f6f7f9;--fg:#171a1f;--mut:#667;--card:#fff;--line:#dde1e7;--pos:#0a7d3b;--neg:#b3261e}}
@media(prefers-color-scheme:dark){{:root{{--bg:#12151a;--fg:#e8eaee;--mut:#98a;--card:#1b1f26;--line:#2c323c;--pos:#4cc07a;--neg:#ff7b72}}}}
body{{font:15px/1.5 system-ui,sans-serif;background:var(--bg);color:var(--fg);margin:0;padding:20px;max-width:1200px;margin-inline:auto}}
h1{{font-size:22px;margin:0 0 4px}} h2{{font-size:17px;margin:28px 0 8px}} h3{{font-size:14px;margin:0 0 6px;color:var(--mut);font-weight:600}}
.muted{{color:var(--mut)}} .grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:12px}}
.stat{{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:14px}} .big{{font-size:30px;font-weight:700}}
.scroll{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;background:var(--card);font-size:13.5px}}
th,td{{padding:6px 10px;border-bottom:1px solid var(--line);text-align:left;white-space:nowrap}} .pos{{color:var(--pos)}} .neg{{color:var(--neg)}}
a{{color:inherit}} nav a{{margin-right:14px}}</style>
<h1>Performance &mdash; {period}</h1>
<p class=muted>{start} to {end}. L1 = booked / last target hit (official). L2 = highest reported (reference).</p>
<nav>{nav}</nav>
<h2>Headline (new signals + carry-ins)</h2><div class=grid>{s1}{s2}{s3}</div>
<h2>New signals only</h2><div class=grid>{s4}{s5}</div>
<h2>Counted trades</h2>{t1}
<h2>Excluded / open</h2>{t2}
<h2>Closed this period but already counted earlier</h2>{t3}
<h2>Needs review ({nrev})</h2>{t4}
<p class=muted><a href="/report.csv?period={period}{k}">CSV</a> &middot; <a href="/report.xlsx?period={period}{k}">Excel</a> &middot; <a href="/review{kq}">parser review queue</a></p>"""


@app.get("/", response_class=HTMLResponse)
def home(period: str = "", key: str = ""):
    need_view(key)
    period = period or current_period()
    r = report_for(period)
    st = r["stats"]
    k = f"&key={key}" if key else ""
    kq = f"?key={key}" if key else ""
    nav = ("".join(f"<a href='/?period={p}{k}'>{p}</a>" for p in
                   (current_period("month"), current_period("week"))))
    unm = [{"symbol": e.symbol, "signal_date": e.event_date, "status": e.kind, "note": "no matching signal"}
           for e in r["unmatched_events"]]
    rev = table(r["review"] + unm, ["signal_date", "symbol", "status", "note"], ["Date", "Stock", "Status", "Note"])
    return PAGE.format(
        period=period, start=r["start"], end=r["end"], nav=nav, k=k, kq=kq,
        s1=stat_block("Incl. provisional", st["all_incl_provisional"]),
        s2=stat_block("Closed only", st["all_closed_only"]),
        s3=stat_block("L2 (reference)", st["l2_all"]),
        s4=stat_block("New, incl. provisional", st["new_incl_provisional"]),
        s5=stat_block("New, closed only", st["new_closed_only"]),
        t1=table(r["rows"]), t2=table(r["excluded"]),
        t3=table(r["closing_already_counted"], COLS[:-2] + ["counted_in", "note"], HEAD[:-2] + ["Counted in", "Note"]),
        t4=rev, nrev=len(r["review"]) + len(unm))


@app.get("/api/report")
def api_report(period: str = "", key: str = ""):
    need_view(key)
    r = report_for(period or current_period())
    return JSONResponse(content=_jsonable(r))


def _jsonable(o):
    if isinstance(o, dict):
        return {k: _jsonable(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_jsonable(v) for v in o]
    if hasattr(o, "isoformat"):
        return o.isoformat()
    if hasattr(o, "__dict__"):
        return _jsonable(vars(o))
    return o


@app.get("/report.csv")
def report_csv(period: str = "", key: str = ""):
    need_view(key)
    r = report_for(period or current_period())
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(HEAD)
    for x in r["rows"] + r["excluded"]:
        w.writerow([x.get(c, "") if x.get(c) is not None else "" for c in COLS])
    return Response(out.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=trades_{r['label']}.csv"})


@app.get("/report.xlsx")
def report_xlsx(period: str = "", key: str = ""):
    need_view(key)
    from openpyxl import Workbook
    r = report_for(period or current_period())
    wb = Workbook()
    ws = wb.active; ws.title = "Counted"
    ws.append(HEAD)
    for x in r["rows"]:
        ws.append([str(x[c]) if c == "signal_date" else x.get(c) for c in COLS])
    ws2 = wb.create_sheet("Excluded"); ws2.append(HEAD)
    for x in r["excluded"]:
        ws2.append([str(x[c]) if c == "signal_date" else x.get(c) for c in COLS])
    ws3 = wb.create_sheet("Stats"); ws3.append(["View", "Counted", "Winners", "Losers", "Win rate %", "Total %", "Avg %"])
    for name, s in r["stats"].items():
        ws3.append([name, s["counted"], s["winners"], s["losers"], s["win_rate"], s["total"], s["avg"]])
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=trades_{r['label']}.xlsx"})


@app.get("/review", response_class=HTMLResponse)
def review(key: str = ""):
    need_view(key)
    q = db.review_queue()
    return "<meta charset=utf-8><h2>Messages the parser could not understand</h2>" + "".join(
        f"<pre>{html.escape(x['posted_at'])}  [{html.escape(x['note'])}]\n{html.escape(x['text'])}</pre>" for x in q) \
        or "<p>Nothing to review.</p>"
